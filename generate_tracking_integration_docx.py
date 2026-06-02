"""
Generate a styled Tracking Integration & Logistics Platform Roadmap (.docx)
for the Fleet Management System project.

Requirements
------------
    pip install python-docx

Usage
-----
    python generate_tracking_integration_docx.py

Output
------
    Tracking_Integration_Roadmap.docx  (created in the current working directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


# ---------------------------------------------------------------------------
# Color palette  (aligned with other docx generators in this repo)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1A, 0x37, 0x6C)
TEAL = RGBColor(0x00, 0x7B, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_GREY = RGBColor(0x40, 0x40, 0x40)


# ---------------------------------------------------------------------------
# Helper utilities  (styling matches generate_sla_docx.py)
# ---------------------------------------------------------------------------

def set_cell_bg(cell, hex_color: str) -> None:
    """Fill a table cell background with a solid color (hex, e.g. '1A376C')."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc: Document) -> None:
    """Insert a thin teal horizontal rule paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "007B83")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(0)


def add_section_heading(doc: Document, number: str, title: str) -> None:
    """Add a numbered section heading (Heading 1 style, navy background strip)."""
    p = doc.add_paragraph(style="Heading 1")
    run = p.add_run(f"  {number}  {title}")
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "1A376C")
    pPr.append(shd)


def add_sub_heading(doc: Document, title: str) -> None:
    """Add a sub-section heading (Heading 2 style, teal text)."""
    p = doc.add_paragraph(style="Heading 2")
    run = p.add_run(title)
    run.font.color.rgb = TEAL
    run.font.bold = True
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str) -> None:
    """Add a normal body paragraph."""
    p = doc.add_paragraph(text, style="Normal")
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.color.rgb = DARK_GREY
        run.font.size = Pt(10.5)


def add_bullet(doc: Document, bold_label: str, body_text: str) -> None:
    """Add a bullet point with an optional bold label followed by body text."""
    p = doc.add_paragraph(style="List Bullet")
    label_run = p.add_run(bold_label)
    label_run.bold = True
    label_run.font.color.rgb = NAVY
    label_run.font.size = Pt(10.5)
    if body_text:
        body_run = p.add_run(" " + body_text)
        body_run.font.color.rgb = DARK_GREY
        body_run.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(4)


def add_styled_table(doc: Document, headers: list, rows: list) -> None:
    """Render a styled data table."""
    num_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=num_cols)
    table.style = "Table Grid"

    hdr_row = table.rows[0]
    for col_idx, hdr_text in enumerate(headers):
        cell = hdr_row.cells[col_idx]
        set_cell_bg(cell, "1A376C")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        run = para.add_run(hdr_text)
        run.bold = True
        run.font.color.rgb = WHITE
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"
        para.paragraph_format.left_indent = Pt(4)

    for row_idx, row_data in enumerate(rows):
        bg = "E8EEF7" if row_idx % 2 == 0 else "F2F2F2"
        t_row = table.rows[1 + row_idx]
        for col_idx, cell_text in enumerate(row_data):
            cell = t_row.cells[col_idx]
            set_cell_bg(cell, bg)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            run = para.add_run(cell_text)
            run.font.color.rgb = DARK_GREY
            run.font.size = Pt(10.5)
            run.font.name = "Calibri"
            para.paragraph_format.left_indent = Pt(4)

    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK_GREY

    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = "Fleet Management System  |  Tracking Integration Roadmap"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_para.runs[0]
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = TEAL
    header_run.font.italic = True

    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text = (
        f"Confidential  ·  Generated {datetime.date.today().strftime('%B %d, %Y')}"
    )
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.runs[0]
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    footer_run.font.italic = True

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Tracking Integration & Logistics Platform Roadmap")
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = NAVY
    title_run.font.name = "Calibri"

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Fleet Management System")
    sub_run.font.size = Pt(16)
    sub_run.font.color.rgb = TEAL
    sub_run.font.name = "Calibri"
    subtitle_p.paragraph_format.space_after = Pt(18)

    add_horizontal_rule(doc)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"
    meta_table.autofit = True

    meta_data = [
        ("Project Name", "Fleet Management System"),
        ("Document Type", "Tracking Integration Roadmap"),
        ("Version", "1.0"),
        ("Date", datetime.date.today().strftime("%B %d, %Y")),
    ]

    for i, (label, value) in enumerate(meta_data):
        row = meta_table.rows[i]

        label_cell = row.cells[0]
        set_cell_bg(label_cell, "1A376C")
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp = label_cell.paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.color.rgb = WHITE
        lr.font.size = Pt(10.5)
        lr.font.name = "Calibri"
        lp.paragraph_format.left_indent = Pt(6)

        value_cell = row.cells[1]
        set_cell_bg(value_cell, "E8EEF7")
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        vp = value_cell.paragraphs[0]
        vr = vp.add_run(value)
        vr.font.color.rgb = DARK_GREY
        vr.font.size = Pt(10.5)
        vr.font.name = "Calibri"
        vp.paragraph_format.left_indent = Pt(6)

    doc.add_paragraph()

    add_section_heading(doc, "1.", "Executive Overview")
    add_body(
        doc,
        "With the Production MVP complete, the next stage focuses on delivering a "
        "tracking-first logistics platform. The roadmap below connects vehicles, "
        "shipments, and telemetry into a single operational view while progressively "
        "unlocking automation, auditing, and performance intelligence.",
    )
    add_sub_heading(doc, "Primary Outcomes")
    add_bullet(
        doc,
        "Real-time visibility:",
        "Give stakeholders an always-current view of shipment movement and status.",
    )
    add_bullet(
        doc,
        "Operational confidence:",
        "Reduce manual status checks and improve exception response times.",
    )
    add_bullet(
        doc,
        "Scalable foundation:",
        "Standardize tracking data to support automation and analytics later phases.",
    )

    add_section_heading(doc, "2.", "Roadmap at a Glance")
    add_body(
        doc,
        "Delivery is staged to prioritize the tracking foundation and live visibility "
        "before expanding into automation, alerts, and analytics.",
    )
    add_styled_table(
        doc,
        headers=["Phase", "Primary Deliverable", "Priority"],
        rows=[
            ("1", "Tracking foundation & integration", "Critical"),
            ("2", "Live shipment visibility", "Critical"),
            ("3", "Route history & playback", "High"),
            ("4", "Geofence automation", "High"),
            ("5", "Status automation engine", "High"),
            ("6", "Notifications & alerting", "Medium"),
            ("7", "Analytics & carrier performance", "Medium"),
        ],
    )

    add_section_heading(doc, "3.", "Phase 1 — Tracking Foundation")
    add_body(
        doc,
        "Establish the technical backbone for tracking by registering vehicles, "
        "linking shipments, and ingesting location updates reliably.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(
        doc,
        "Vehicle registry:",
        "Maintain vehicle profiles, identifiers, and carrier ownership details.",
    )
    add_bullet(
        doc,
        "Shipment assignment:",
        "Bind shipments to vehicles with active and historical relationships.",
    )
    add_bullet(
        doc,
        "Telemetry ingestion:",
        "Receive position updates, status signals, and last-seen timestamps.",
    )
    add_bullet(
        doc,
        "Administration:",
        "Manage registrations, devices, and tracking configuration settings.",
    )
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Vehicle management module:", "Vehicle profiles and carrier links.")
    add_bullet(doc, "Shipment-to-vehicle assignment:", "Assignment workflow and history.")
    add_bullet(doc, "Tracking integration layer:", "Location processing and storage.")
    add_bullet(doc, "Tracking admin console:", "Configuration and device management.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Foundational readiness:", "Enables all downstream tracking features.")
    add_bullet(doc, "Better oversight:", "Improves visibility into active operations.")

    add_section_heading(doc, "4.", "Phase 2 — Live Shipment Tracking")
    add_body(
        doc,
        "Expose real-time tracking views for shippers, carriers, and administrators "
        "with map-based monitoring and up-to-date status indicators.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(doc, "Live location:", "Show current vehicle position with last update time.")
    add_bullet(doc, "Tracking dashboard:", "Monitor active shipments on a real-time map.")
    add_bullet(
        doc,
        "Shipment tracking page:",
        "Display shipment, carrier, vehicle, location, status, and activity details.",
    )
    add_bullet(
        doc,
        "History capture:",
        "Store position history for later review and investigation.",
    )
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Live tracking dashboard:", "Operational map with active shipments.")
    add_bullet(doc, "Shipment tracking interface:", "Dedicated shipment visibility page.")
    add_bullet(doc, "Historical tracking records:", "Queryable movement history.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Customer confidence:", "Fewer status calls and higher transparency.")
    add_bullet(doc, "Operational control:", "Immediate awareness of shipment progress.")

    add_section_heading(doc, "5.", "Phase 3 — Route History & Playback")
    add_body(
        doc,
        "Enable route playback and journey analysis for completed and in-progress "
        "shipments to support audits and operational review.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(doc, "Route playback:", "Replay historical movement on a timeline.")
    add_bullet(
        doc,
        "Trip analytics:",
        "Summarize distance traveled, stops, and travel behavior.",
    )
    add_bullet(
        doc,
        "Investigation tools:",
        "Validate deliveries and verify compliance with expected routes.",
    )
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Playback interface:", "Time-based route visualization.")
    add_bullet(doc, "Journey reports:", "Exportable movement summaries.")
    add_bullet(doc, "Route analysis tools:", "Stop duration and distance insights.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Audit readiness:", "Improves dispute resolution and verification.")
    add_bullet(doc, "Performance insights:", "Highlights operational efficiency gaps.")

    add_section_heading(doc, "6.", "Phase 4 — Geofence Automation")
    add_body(
        doc,
        "Introduce location-based triggers for pickup and delivery sites to automate "
        "arrival/departure events and build a reliable audit trail.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(doc, "Pickup geofences:", "Detect arrival and departure at origin sites.")
    add_bullet(doc, "Destination geofences:", "Confirm delivery zone entry and exit.")
    add_bullet(doc, "Event monitoring:", "Track entry, exit, and arrival milestones.")
    add_bullet(doc, "Event logging:", "Persist geofence activity for audits.")
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Geofence management:", "Create and manage location boundaries.")
    add_bullet(doc, "Event processing:", "Real-time geofence trigger handling.")
    add_bullet(doc, "Automated event history:", "Structured geofence audit trail.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Reduced manual checks:", "Automates pickup/delivery confirmation.")
    add_bullet(doc, "Higher accuracy:", "Improves delivery verification quality.")

    add_section_heading(doc, "7.", "Phase 5 — Intelligent Status Automation")
    add_body(
        doc,
        "Automate shipment status progression using tracking events and exceptions "
        "to minimize manual updates and improve data accuracy.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(
        doc,
        "Event-driven progression:",
        "Advance statuses based on arrivals, departures, and delivery completion.",
    )
    add_bullet(
        doc,
        "Workflow automation:",
        "Reduce user intervention with rules-based status updates.",
    )
    add_bullet(
        doc,
        "Exception detection:",
        "Identify delays, route deviations, and unexpected stops.",
    )
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Status automation engine:", "Rules-based shipment progression.")
    add_bullet(doc, "Event processing framework:", "Unified event handling pipeline.")
    add_bullet(doc, "Automation rule library:", "Configurable status triggers.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Operational efficiency:", "Less manual effort and faster updates.")
    add_bullet(doc, "Data accuracy:", "Consistent status and milestone reporting.")

    add_section_heading(doc, "8.", "Phase 6 — Notifications & Alerts")
    add_body(
        doc,
        "Keep stakeholders informed with proactive shipment notifications and "
        "location-based alerts for critical tracking events.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(doc, "Shipment notifications:", "Assignment, pickup, transit, delivery updates.")
    add_bullet(
        doc,
        "Location alerts:",
        "Arrival and departure notifications for pickup/destination geofences.",
    )
    add_bullet(
        doc,
        "Administrative alerts:",
        "Detect tracking interruptions and delay conditions.",
    )
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Notification service:", "Multi-channel update engine.")
    add_bullet(doc, "Alert management:", "Configurable alert rules and routing.")
    add_bullet(doc, "User notification center:", "Centralized notification history.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Faster communication:", "Improves stakeholder awareness.")
    add_bullet(doc, "Reduced delays:", "Early warning for operational issues.")

    add_section_heading(doc, "9.", "Phase 7 — Analytics & Carrier Performance")
    add_body(
        doc,
        "Use tracking data to measure carrier performance, vehicle utilization, "
        "and service reliability through dashboards and ratings.",
    )
    add_sub_heading(doc, "Core Capabilities")
    add_bullet(
        doc,
        "Performance metrics:",
        "On-time delivery, completion rates, and average transit times.",
    )
    add_bullet(
        doc,
        "Utilization insights:",
        "Vehicle operating hours, distance traveled, and asset usage.",
    )
    add_bullet(
        doc,
        "Operational dashboards:",
        "Fleet and shipment activity reporting for leadership.",
    )
    add_bullet(
        doc,
        "Ratings & reviews:",
        "Capture post-delivery feedback and carrier scores.",
    )
    add_sub_heading(doc, "Deliverables")
    add_bullet(doc, "Performance dashboards:", "Carrier and fleet analytics.")
    add_bullet(doc, "Rating system:", "Post-delivery feedback integration.")
    add_bullet(doc, "Operational reports:", "Exportable performance insights.")
    add_sub_heading(doc, "Business Value")
    add_bullet(doc, "Data-driven decisions:", "Enables accountability and optimization.")
    add_bullet(doc, "Service quality visibility:", "Improves trust and marketplace health.")

    add_section_heading(doc, "10.", "Strategic Outcome")
    add_body(
        doc,
        "This phased approach transforms the platform from shipment management into a "
        "fully trackable logistics ecosystem, enabling real-time visibility, verified "
        "delivery events, automated workflows, and actionable performance insights.",
    )

    return doc


if __name__ == "__main__":
    document = build_document()
    document.save("Tracking_Integration_Roadmap.docx")
