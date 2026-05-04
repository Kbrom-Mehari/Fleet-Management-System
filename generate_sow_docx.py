"""
Generate a well-designed Statement of Work (SoW) .docx document
for the Fleet Management System project.

Requirements
------------
    pip install python-docx

Usage
-----
    python generate_sow_docx.py

Output
------
    Statement_of_Work.docx  (created in the current working directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
NAVY       = RGBColor(0x1A, 0x37, 0x6C)   # primary heading colour
TEAL       = RGBColor(0x00, 0x7B, 0x83)   # secondary / accent
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_GREY  = RGBColor(0x40, 0x40, 0x40)


# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def set_cell_bg(cell, hex_color: str) -> None:
    """Fill a table cell background with a solid colour (hex, e.g. '1A376C')."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc: Document) -> None:
    """Insert a thin horizontal rule paragraph."""
    p      = doc.add_paragraph()
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "007B83")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(0)


def add_section_heading(doc: Document, number: str, title: str) -> None:
    """Add a numbered section heading (Heading 1 style, navy background strip)."""
    p    = doc.add_paragraph(style="Heading 1")
    run  = p.add_run(f"  {number}  {title}")
    run.font.color.rgb = WHITE
    run.font.bold      = True
    run.font.size      = Pt(13)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)
    # Shade the paragraph background
    pPr  = p._p.get_or_add_pPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "1A376C")
    pPr.append(shd)


def add_sub_heading(doc: Document, title: str) -> None:
    """Add a sub-section heading (Heading 2 style, teal text)."""
    p   = doc.add_paragraph(style="Heading 2")
    run = p.add_run(title)
    run.font.color.rgb = TEAL
    run.font.bold      = True
    run.font.size      = Pt(11)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)


def add_body(doc: Document, text: str) -> None:
    """Add a normal body paragraph."""
    p = doc.add_paragraph(text, style="Normal")
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.color.rgb = DARK_GREY
        run.font.size      = Pt(10.5)


def add_bullet(doc: Document, bold_label: str, body_text: str) -> None:
    """Add a bullet point with an optional bold label followed by body text."""
    p         = doc.add_paragraph(style="List Bullet")
    label_run = p.add_run(bold_label)
    label_run.bold           = True
    label_run.font.color.rgb = NAVY
    label_run.font.size      = Pt(10.5)
    if body_text:
        body_run               = p.add_run(" " + body_text)
        body_run.font.color.rgb = DARK_GREY
        body_run.font.size      = Pt(10.5)
    p.paragraph_format.space_after = Pt(4)


def add_phase_heading(doc: Document, phase: str) -> None:
    """Add a phase heading inside Section 4 (teal, italic)."""
    p   = doc.add_paragraph(style="Normal")
    run = p.add_run(phase)
    run.bold             = True
    run.italic           = True
    run.font.color.rgb   = TEAL
    run.font.size        = Pt(11)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.54)
        section.right_margin  = Cm(2.54)

    # ── Default Normal style ─────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name    = "Calibri"
    normal.font.size    = Pt(10.5)
    normal.font.color.rgb = DARK_GREY

    # ── Header ───────────────────────────────────────────────────────────────
    header      = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text      = "Fleet Management System  |  Statement of Work"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run            = header_para.runs[0]
    header_run.font.size      = Pt(9)
    header_run.font.color.rgb = TEAL
    header_run.font.italic    = True

    # ── Footer ───────────────────────────────────────────────────────────────
    footer      = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text      = f"Confidential  ·  Generated {datetime.date.today().strftime('%B %d, %Y')}"
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run            = footer_para.runs[0]
    footer_run.font.size      = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    footer_run.font.italic    = True

    # ════════════════════════════════════════════════════════════════════════
    # COVER / TITLE BLOCK
    # ════════════════════════════════════════════════════════════════════════
    # Title paragraph
    title_p       = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run         = title_p.add_run("Statement of Work")
    title_run.bold           = True
    title_run.font.size      = Pt(26)
    title_run.font.color.rgb = NAVY
    title_run.font.name      = "Calibri"

    subtitle_p   = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run          = subtitle_p.add_run("& Project Delivery Document")
    sub_run.font.size      = Pt(16)
    sub_run.font.color.rgb = TEAL
    sub_run.font.name      = "Calibri"
    subtitle_p.paragraph_format.space_after = Pt(18)

    add_horizontal_rule(doc)

    # ── Metadata table ───────────────────────────────────────────────────────
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.style = "Table Grid"
    meta_table.autofit = True

    meta_data = [
        ("Project Name",    "Fleet Management System"),
        ("Platform Core",   "Traccar-based Architecture"),
        ("Date",            "May 4, 2026"),
    ]

    for i, (label, value) in enumerate(meta_data):
        row = meta_table.rows[i]

        # Label cell
        label_cell = row.cells[0]
        set_cell_bg(label_cell, "1A376C")
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp  = label_cell.paragraphs[0]
        lr  = lp.add_run(label)
        lr.bold           = True
        lr.font.color.rgb = WHITE
        lr.font.size      = Pt(10.5)
        lr.font.name      = "Calibri"
        lp.paragraph_format.left_indent = Pt(6)

        # Value cell
        value_cell = row.cells[1]
        set_cell_bg(value_cell, "E8EEF7")
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        vp  = value_cell.paragraphs[0]
        vr  = vp.add_run(value)
        vr.font.color.rgb = DARK_GREY
        vr.font.size      = Pt(10.5)
        vr.font.name      = "Calibri"
        vp.paragraph_format.left_indent = Pt(6)

    doc.add_paragraph()   # spacer

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 1 – Overview
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "1.", "Overview of What Was Built and Delivered")

    add_body(doc,
        "The project involved the design, development, and deployment of a comprehensive, "
        "real-time Fleet Management and GPS Tracking System. Built upon the robust and "
        "scalable Traccar engine, the platform is engineered to provide end-to-end "
        "visibility into fleet operations."
    )
    add_body(doc,
        "The delivered system empowers administrators and fleet managers to monitor vehicle "
        "locations in real-time, enforce security protocols, analyse operational efficiency, "
        "and generate actionable insights through advanced reporting."
    )
    add_body(doc,
        "The final deliverable includes a fully functional web-based dashboard, backend "
        "server infrastructure configured for high-frequency GPS data ingestion, and "
        "integrated support for a wide range of GPS tracking devices."
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 2 – Features
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "2.", "Features and Capabilities Delivered")

    add_body(doc,
        "The delivered platform encompasses all native capabilities of the Traccar engine, "
        "enhanced and configured to meet the project's operational requirements. "
        "Key features include:"
    )

    features = [
        (
            "Real-Time GPS Tracking:",
            "Live mapping and continuous location updates for all registered fleet vehicles "
            "with high accuracy and low latency.",
        ),
        (
            "Remote Immobilisation & Security:",
            "Remote engine cut-off and restoration capabilities to prevent theft and enforce "
            "unauthorised-usage policies directly from the dashboard.",
        ),
        (
            "Fuel Control & Monitoring:",
            "Advanced fuel-level tracking, consumption analytics, and automated alerts for "
            "sudden drops (fuel theft) or abnormal usage patterns.",
        ),
        (
            "Advanced Reports & Analytics:",
            "Comprehensive data-analysis tools including visual charts and graphs. Reports "
            "can be generated for trips, stops, summary metrics, and events, with seamless "
            "export capabilities to PDF and CSV formats.",
        ),
        (
            "Geofencing:",
            "Creation of custom virtual perimeters (polygons, circles, or lines) with "
            "automated alerts triggered upon vehicle entry or exit.",
        ),
        (
            "Travel History & Route Playback:",
            "Stored historical tracking data allowing managers to replay past routes, "
            "analyse driving patterns, and verify completed trips.",
        ),
        (
            "Alerts & Notifications:",
            "Configurable real-time notifications (via web, email, or SMS/push integrations) "
            "for events such as overspeeding, harsh braking, device tampering, and offline "
            "statuses.",
        ),
        (
            "Device & User Management:",
            "Hierarchical user access controls allowing administrators to group devices, "
            "assign specific vehicles to sub-users, and restrict access based on roles.",
        ),
    ]

    for label, body in features:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 3 – QA Summary
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "3.", "Quality Assurance Summary")

    add_body(doc,
        "Prior to deployment, the system underwent rigorous Quality Assurance (QA) testing "
        "to ensure reliability, security, and performance under operational loads:"
    )

    qa_items = [
        (
            "Functional Testing:",
            "All core features (tracking, geofencing, reports, and immobilisation commands) "
            "were successfully validated against expected behaviours using simulated and "
            "physical GPS hardware.",
        ),
        (
            "Performance & Load Testing:",
            "The data-ingestion pipeline was tested to handle concurrent connections and "
            "high-frequency coordinate updates without data loss or dashboard latency.",
        ),
        (
            "Security Testing:",
            "User authentication, role-based access controls, and secure data-transmission "
            "protocols (SSL/TLS) were verified to protect sensitive fleet data and prevent "
            "unauthorised vehicle control.",
        ),
        (
            "Cross-Browser & Device Compatibility:",
            "The web interface was tested across major browsers (Chrome, Firefox, Safari, "
            "Edge) and screen sizes to ensure a responsive and accessible user experience.",
        ),
    ]

    for label, body in qa_items:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 4 – Deployment & Handover
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "4.", "Deployment and Handover Plan")

    add_body(doc,
        "To ensure a smooth transition to the operational team, the deployment and handover "
        "will be executed in the following phases:"
    )

    # Phase 1
    add_phase_heading(doc, "Phase 1: Environment Setup & Deployment")
    for item in [
        "Provisioning and configuration of the production server environment.",
        "Installation of the Traccar backend, database setup (MySQL/PostgreSQL), and web-interface deployment.",
        "Configuration of network ports and firewalls to securely accept incoming GPS device protocols.",
        "SSL certificate installation for secure web access.",
    ]:
        add_bullet(doc, "", item)

    # Phase 2
    add_phase_heading(doc, "Phase 2: System Configuration & Onboarding")
    for item in [
        "Registration of initial admin accounts and setup of role-based user groups.",
        "Configuration of baseline system settings including regional timezones, map layers, and default notification templates.",
        "Onboarding of the initial batch of GPS tracking devices to verify live production connectivity.",
    ]:
        add_bullet(doc, "", item)

    # Phase 3
    add_phase_heading(doc, "Phase 3: Handover & Training")
    for label, body in [
        (
            "Documentation:",
            "Delivery of system documentation including Administrator Guides (for managing "
            "devices, users, and server settings) and User Manuals (for daily tracking and "
            "reporting).",
        ),
        (
            "Training Sessions:",
            "Scheduled walkthroughs with fleet managers covering real-time tracking, report "
            "generation, setting up geofences, and executing remote immobilisation.",
        ),
        (
            "Asset Transfer:",
            "Handover of all server credentials, database access, and administrative "
            "accounts to the primary stakeholder.",
        ),
    ]:
        add_bullet(doc, label, body)

    # Phase 4
    add_phase_heading(doc, "Phase 4: Post-Deployment Support")
    add_bullet(doc, "Hypercare Period:",
        "A designated 14–30 day support window to address minor bugs, assist with device "
        "protocol troubleshooting, and provide operational guidance as users acclimate to "
        "the new system."
    )

    # ── Final rule ────────────────────────────────────────────────────────────
    doc.add_paragraph()
    add_horizontal_rule(doc)
    closing   = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = closing.add_run("End of Statement of Work Document")
    cr.italic           = True
    cr.font.size        = Pt(9)
    cr.font.color.rgb   = RGBColor(0x88, 0x88, 0x88)

    return doc


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    output_path = "Statement_of_Work.docx"
    doc = build_document()
    doc.save(output_path)
    print(f"✅  Document saved: {output_path}")
