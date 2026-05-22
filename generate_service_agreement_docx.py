"""
Generate a simple Service Agreement document (.docx)
for the Fleet Management System project.

Requirements
------------
    pip install python-docx

Usage
-----
    python generate_service_agreement_docx.py

Output
------
    Simple_Service_Agreement.docx  (created in the current working directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime


NAVY = RGBColor(0x1A, 0x37, 0x6C)
TEAL = RGBColor(0x00, 0x7B, 0x83)
DARK_GREY = RGBColor(0x40, 0x40, 0x40)


def set_document_defaults(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK_GREY

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = NAVY

    heading_1 = doc.styles["Heading 1"]
    heading_1.font.name = "Calibri"
    heading_1.font.size = Pt(13)
    heading_1.font.bold = True
    heading_1.font.color.rgb = NAVY

    heading_2 = doc.styles["Heading 2"]
    heading_2.font.name = "Calibri"
    heading_2.font.size = Pt(11)
    heading_2.font.bold = True
    heading_2.font.color.rgb = TEAL


def add_header_footer(doc: Document) -> None:
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = "Fleet Management System  |  Simple Service Agreement"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_para.runs[0]
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = TEAL
    header_run.font.italic = True

    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text = f"Confidential  ·  Generated {datetime.date.today().strftime('%B %d, %Y')}"
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.runs[0]
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    footer_run.font.italic = True


def add_section(doc: Document, title: str) -> None:
    doc.add_paragraph(title, style="Heading 1")


def build_document() -> Document:
    doc = Document()
    set_document_defaults(doc)
    add_header_footer(doc)

    doc.add_paragraph("Simple Service Agreement", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "This Service Agreement is made between the Service Provider and the Client as of "
        "the Effective Date written below.",
        style="Normal",
    )

    add_section(doc, "1. Parties")
    doc.add_paragraph("Service Provider: Kbrom Mehari", style="List Bullet")
    doc.add_paragraph("Client: Mehari Hiwot", style="List Bullet")
    doc.add_paragraph("Service Provider Address/Contact: [Insert address, phone, email]", style="List Bullet")
    doc.add_paragraph("Client Address/Contact: [Insert address, phone, email]", style="List Bullet")

    add_section(doc, "2. Scope of Services")
    doc.add_paragraph(
        "The Service Provider will deliver a fleet tracking platform, deploy it in the agreed "
        "environment, and complete handover to the Client.",
        style="Normal",
    )
    doc.add_paragraph("Platform delivery and configuration", style="List Bullet")
    doc.add_paragraph("Deployment to production environment", style="List Bullet")
    doc.add_paragraph("Admin account handover and basic orientation", style="List Bullet")

    add_section(doc, "3. Pricing & Payment Terms")
    doc.add_paragraph("Total Price: 250,000 ETB", style="Normal")

    doc.add_paragraph("Payment Schedule", style="Heading 2")
    doc.add_paragraph("50% (125,000 ETB) due upon signing of the agreement.", style="List Bullet")
    doc.add_paragraph(
        "50% (125,000 ETB) due after deployment is completed and the Client has received the admin account (handover).",
        style="List Bullet",
    )

    doc.add_paragraph("Accepted Payment Methods", style="Heading 2")
    doc.add_paragraph("Bank transfer: [Insert bank/account details]", style="List Bullet")
    doc.add_paragraph("Mobile money: [Insert provider/account details]", style="List Bullet")
    doc.add_paragraph("Cash: [Insert process/receipt details]", style="List Bullet")

    doc.add_paragraph("Late Payment Policy", style="Heading 2")
    doc.add_paragraph("Late fee: [X%] per month on overdue amounts after [Y] days.", style="List Bullet")

    add_section(doc, "4. Term & Termination")
    doc.add_paragraph(
        "Term: This Agreement starts on [Effective Date] and continues until delivery and handover are completed, "
        "unless terminated earlier as set out below.",
        style="Normal",
    )
    doc.add_paragraph("Either party may terminate with [X] days written notice if the other party materially breaches this Agreement.", style="List Bullet")
    doc.add_paragraph("Client shall pay for all completed work and approved costs up to the termination date.", style="List Bullet")

    add_section(doc, "5. Governing Law")
    doc.add_paragraph("Ethiopia", style="List Bullet")

    add_section(doc, "6. Signatures")
    doc.add_paragraph("Service Provider: Kbrom Mehari", style="Normal")
    doc.add_paragraph("Signature: ____________________________", style="Normal")
    doc.add_paragraph("Date: ________________________________", style="Normal")
    doc.add_paragraph("", style="Normal")
    doc.add_paragraph("Client: Mehari Hiwot", style="Normal")
    doc.add_paragraph("Signature: ____________________________", style="Normal")
    doc.add_paragraph("Date: ________________________________", style="Normal")

    return doc


if __name__ == "__main__":
    output_path = "Simple_Service_Agreement.docx"
    document = build_document()
    document.save(output_path)
    print(f"✅  Document saved: {output_path}")
