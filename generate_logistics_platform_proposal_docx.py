"""
Generate a polished Logistics Platform Proposal (.docx).

Requirements
------------
    pip install python-docx

Usage
-----
    python generate_logistics_platform_proposal_docx.py

Output
------
    Logistics_Platform_Proposal.docx  (created in the current working directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1A, 0x37, 0x6C)      # primary heading colour
TEAL = RGBColor(0x00, 0x7B, 0x83)      # secondary / accent
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_GREY = RGBColor(0x40, 0x40, 0x40)


# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def set_cell_bg(cell, hex_color: str) -> None:
    """Fill a table cell background with a solid colour (hex, e.g. '1A376C')."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc: Document) -> None:
    """Insert a thin teal horizontal rule paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "007B83")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(0)


def add_section_heading(doc: Document, number: str, title: str) -> None:
    """Add a numbered section heading (Heading 1 style, navy background strip)."""
    p = doc.add_paragraph(style="Heading 1")
    run = p.add_run(f"  {number}  {title}")
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "1A376C")
    pPr.append(shd)


def add_sub_heading(doc: Document, title: str) -> None:
    """Add a sub-section heading (Heading 2 style, teal text)."""
    p = doc.add_paragraph(style="Heading 2")
    run = p.add_run(title)
    run.font.color.rgb = TEAL
    run.font.bold = True
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str) -> None:
    """Add a normal body paragraph."""
    p = doc.add_paragraph(text, style="Normal")
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.color.rgb = DARK_GREY
        run.font.size = Pt(10.5)


def add_bullet(doc: Document, bold_label: str, body_text: str) -> None:
    """Add a bullet point with an optional bold label followed by body text."""
    p = doc.add_paragraph(style="List Bullet")
    label_run = p.add_run(bold_label)
    label_run.bold = True
    label_run.font.color.rgb = NAVY
    label_run.font.size = Pt(10.5)
    if body_text:
        body_run = p.add_run(" " + body_text)
        body_run.font.color.rgb = DARK_GREY
        body_run.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(4)


def add_highlight_box(doc: Document, title: str, items: list[str]) -> None:
    """Add a small shaded highlight box with a title and bullet list."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, "E8EEF7")
    para = cell.paragraphs[0]
    run = para.add_run(title)
    run.bold = True
    run.font.color.rgb = NAVY
    run.font.size = Pt(11)
    for item in items:
        p = cell.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.color.rgb = DARK_GREY
        r.font.size = Pt(10.5)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # ── Default Normal style ─────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK_GREY

    # ── Header ───────────────────────────────────────────────────────────────
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = "Logistics Platform  |  Proposal"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_para.runs[0]
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = TEAL
    header_run.font.italic = True

    # ── Footer ───────────────────────────────────────────────────────────────
    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text = f"Confidential  ·  Generated {datetime.date.today().strftime('%B %d, %Y')}"
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.runs[0]
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    footer_run.font.italic = True

    # ════════════════════════════════════════════════════════════════════════
    # COVER / TITLE BLOCK
    # ════════════════════════════════════════════════════════════════════════
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Logistics Platform Proposal")
    title_run.bold = True
    title_run.font.size = Pt(26)
    title_run.font.color.rgb = NAVY
    title_run.font.name = "Calibri"

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Digital Workflows • Real-Time Visibility • Optimized Routes")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = TEAL
    sub_run.font.name = "Calibri"
    subtitle_p.paragraph_format.space_after = Pt(18)

    add_horizontal_rule(doc)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"
    meta_table.autofit = True
    meta_data = [
        ("Prepared For", "Client Organization"),
        ("Prepared By", "Fleet Management Solutions Team"),
        ("Document Type", "Logistics Platform Proposal"),
        ("Date", datetime.date.today().strftime("%B %d, %Y")),
    ]

    for i, (label, value) in enumerate(meta_data):
        row = meta_table.rows[i]
        label_cell = row.cells[0]
        set_cell_bg(label_cell, "1A376C")
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp = label_cell.paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.color.rgb = WHITE
        lr.font.size = Pt(10.5)
        lr.font.name = "Calibri"
        lp.paragraph_format.left_indent = Pt(6)

        value_cell = row.cells[1]
        set_cell_bg(value_cell, "E8EEF7")
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        vp = value_cell.paragraphs[0]
        vr = vp.add_run(value)
        vr.font.color.rgb = DARK_GREY
        vr.font.size = Pt(10.5)
        vr.font.name = "Calibri"
        vp.paragraph_format.left_indent = Pt(6)

    doc.add_paragraph()

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 1 – Executive Summary
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "1.", "Executive Summary")
    add_body(
        doc,
        "This proposal presents a modern logistics platform that digitizes traditional workflows "
        "from shipment creation through delivery, settlement, and performance review. The solution "
        "connects shippers, carriers, and operations teams in one coordinated environment that "
        "reduces manual coordination and improves service reliability."
    )
    add_body(
        doc,
        "The platform prioritizes real-time live shipment tracking and intelligent route optimization "
        "to deliver accurate ETAs, proactive exception management, and superior customer experience."
    )

    add_highlight_box(
        doc,
        "Business Outcomes",
        [
            "Cut dispatch time with smart shipment creation and automated carrier matching.",
            "Improve on-time performance using live tracking and predictive ETAs.",
            "Reduce cost per shipment by minimizing empty miles and route deviations.",
            "Increase shipper confidence with transparent status events and proof of delivery.",
        ],
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 2 – Solution Overview
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "2.", "Solution Overview")
    add_body(
        doc,
        "The platform unifies shipment planning, carrier selection, execution, and settlement in a "
        "single digital workflow. It supports multi-tenant operations, configurable approval rules, "
        "and API-ready integration for ERP, TMS, and financial systems."
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 3 – Core Workflow Modules
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "3.", "Core Workflow Modules")

    add_sub_heading(doc, "3.1 Smart Shipment Creation & Carrier Matching")
    add_body(
        doc,
        "Create shipments in seconds with structured templates, service-level rules, and automatic "
        "carrier matching based on lanes, capacity, price, and performance history."
    )
    for label, body in [
        ("Guided shipment forms:", "Validated fields, commodity rules, and document capture."),
        ("Posting & tendering:", "Broadcast to qualified carriers with configurable acceptance flows."),
        ("Carrier marketplace:", "Rate cards, spot bids, and smart recommendations."),
        ("Approval workflows:", "Enforce budget limits and service SLAs before confirmation."),
    ]:
        add_bullet(doc, label, body)

    add_sub_heading(doc, "3.2 Status & Event Tracking")
    add_body(
        doc,
        "Every shipment is monitored through milestones, exceptions, and automated status events "
        "to keep teams and customers aligned."
    )
    for label, body in [
        ("Milestone updates:", "Pickup, in-transit, at-hub, and delivery confirmations."),
        ("Exception alerts:", "Delays, route deviations, temperature issues, or dwell-time flags."),
        ("Customer notifications:", "Email/SMS updates with configurable frequency."),
        ("Document management:", "Proof of delivery, signatures, and shipping documents stored centrally."),
    ]:
        add_bullet(doc, label, body)

    add_sub_heading(doc, "3.3 Settlement & Payment")
    add_body(
        doc,
        "Integrated payment workflows simplify invoicing, reconciliation, and carrier settlement "
        "while improving cashflow visibility."
    )
    for label, body in [
        ("Automated invoicing:", "Generate invoices from completed shipments and rate cards."),
        ("Dispute handling:", "Track discrepancies with audit trails and approval history."),
        ("Carrier payouts:", "Scheduled settlements with payment status visibility."),
        ("Revenue dashboards:", "Margin, cost per mile, and profitability analytics."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 4 – Real-Time Live Shipment Tracking
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "4.", "Real-Time Live Shipment Tracking")
    add_body(
        doc,
        "Live tracking is the centerpiece of the platform. It provides continuous, map-based "
        "visibility for every shipment, enabling teams to intervene early, reassure customers, "
        "and protect service-level commitments."
    )
    for label, body in [
        ("Instant visibility:", "Second-by-second location updates from carrier devices and mobile apps."),
        ("Proactive exception response:", "Trigger alerts the moment a shipment is off-route or delayed."),
        ("Trusted ETAs:", "Live feeds refine arrival estimates to reduce missed deliveries."),
        ("Customer confidence:", "Shareable tracking links reduce inbound status calls."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 5 – Route Optimization & ETA Intelligence
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "5.", "Route Optimization & ETA Intelligence")
    add_body(
        doc,
        "Route optimization turns operational data into direct cost savings. The platform analyzes "
        "distance, traffic, stop sequences, and delivery constraints to recommend the best route and "
        "continuously update ETAs as conditions change."
    )
    for label, body in [
        ("Lower operating costs:", "Reduce fuel spend and empty miles with optimized routing."),
        ("On-time delivery gains:", "Prioritized stop sequencing and dynamic re-routing."),
        ("Accurate planning:", "ETA predictions update with live conditions and driver status."),
        ("Better asset utilization:", "Balance workloads and maximize fleet productivity."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 6 – Trust, Reviews & Quality Control
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "6.", "Trust, Reviews & Quality Control")
    add_body(
        doc,
        "A review and rating system builds accountability and improves carrier performance over time."
    )
    for label, body in [
        ("Post-delivery ratings:", "Shippers rate carriers and drivers on service quality."),
        ("Performance scorecards:", "On-time %, damage rate, and compliance metrics."),
        ("Carrier qualification:", "Automated eligibility rules based on historical performance."),
        ("Issue resolution:", "Structured feedback loops for continuous improvement."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 7 – Additional Capabilities
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "7.", "Additional Capabilities")
    add_body(
        doc,
        "Supporting features ensure the platform is easy to adopt and scales with operational needs."
    )
    for label, body in [
        ("Role-based access:", "Granular permissions for shippers, carriers, and admins."),
        ("Analytics dashboards:", "Shipment volume, cost trends, and service KPIs."),
        ("Integration APIs:", "Connect ERP, WMS, and accounting systems securely."),
        ("Multi-branch operations:", "Support regional offices and partner networks."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 8 – Security & Reliability
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "8.", "Security & Reliability")
    add_body(
        doc,
        "Security, uptime, and data integrity are built into every layer of the platform."
    )
    for label, body in [
        ("Encrypted communications:", "TLS-secured APIs and device connections."),
        ("Audit logging:", "Track every user action and shipment update."),
        ("High availability:", "Redundant services and automated backups."),
        ("Data residency options:", "Deploy on-premise or in approved cloud regions."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 9 – Implementation & Support
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "9.", "Implementation & Support")
    add_body(
        doc,
        "We deliver a structured rollout plan focused on adoption, performance, and measurable ROI."
    )
    for label, body in [
        ("Discovery workshops:", "Confirm workflows, lanes, and carrier network requirements."),
        ("Configuration & onboarding:", "Set up workflows, roles, and integrations."),
        ("Pilot launch:", "Validate real-time tracking and route optimization impact."),
        ("Ongoing success:", "Continuous optimization and feature enhancements."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 10 – Why This Platform
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "10.", "Why This Platform")
    add_body(
        doc,
        "This logistics platform delivers a measurable uplift in visibility, cost control, and service "
        "quality. By combining real-time tracking with route intelligence, teams gain the confidence "
        "to promise tighter delivery windows and consistently meet them."
    )
    add_body(
        doc,
        "We welcome the opportunity to tailor the platform to your operational model and launch a "
        "pilot that proves value quickly before scaling to full production."
    )

    doc.add_paragraph()
    add_horizontal_rule(doc)
    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = closing.add_run("End of Proposal")
    cr.italic = True
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    return doc


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    output_path = "Logistics_Platform_Proposal.docx"
    doc = build_document()
    doc.save(output_path)
    print(f"✅  Document saved: {output_path}")
