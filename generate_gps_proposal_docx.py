"""
Generate a well-designed GPS Tracking Software Proposal (.docx).

Requirements
------------
    pip install python-docx

Usage
-----
    python generate_gps_proposal_docx.py

Output
------
    GPS_Tracking_Proposal.docx  (created in the current working directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
NAVY       = RGBColor(0x1A, 0x37, 0x6C)   # primary heading colour
TEAL       = RGBColor(0x00, 0x7B, 0x83)   # secondary / accent
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_GREY  = RGBColor(0x40, 0x40, 0x40)


# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def set_cell_bg(cell, hex_color: str) -> None:
    """Fill a table cell background with a solid colour (hex, e.g. '1A376C')."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc: Document) -> None:
    """Insert a thin teal horizontal rule paragraph."""
    p      = doc.add_paragraph()
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "007B83")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(0)


def add_section_heading(doc: Document, number: str, title: str) -> None:
    """Add a numbered section heading (Heading 1 style, navy background strip)."""
    p   = doc.add_paragraph(style="Heading 1")
    run = p.add_run(f"  {number}  {title}")
    run.font.color.rgb = WHITE
    run.font.bold      = True
    run.font.size      = Pt(13)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "1A376C")
    pPr.append(shd)


def add_sub_heading(doc: Document, title: str) -> None:
    """Add a sub-section heading (Heading 2 style, teal text)."""
    p   = doc.add_paragraph(style="Heading 2")
    run = p.add_run(title)
    run.font.color.rgb = TEAL
    run.font.bold      = True
    run.font.size      = Pt(11)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)


def add_body(doc: Document, text: str) -> None:
    """Add a normal body paragraph."""
    p = doc.add_paragraph(text, style="Normal")
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.color.rgb = DARK_GREY
        run.font.size      = Pt(10.5)


def add_bullet(doc: Document, bold_label: str, body_text: str) -> None:
    """Add a bullet point with an optional bold label followed by body text."""
    p         = doc.add_paragraph(style="List Bullet")
    label_run = p.add_run(bold_label)
    label_run.bold           = True
    label_run.font.color.rgb = NAVY
    label_run.font.size      = Pt(10.5)
    if body_text:
        body_run               = p.add_run(" " + body_text)
        body_run.font.color.rgb = DARK_GREY
        body_run.font.size      = Pt(10.5)
    p.paragraph_format.space_after = Pt(4)


def add_highlight_box(doc: Document, title: str, items: list[str]) -> None:
    """Add a small shaded highlight box with a title and bullet list."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, "E8EEF7")
    para = cell.paragraphs[0]
    run = para.add_run(title)
    run.bold = True
    run.font.color.rgb = NAVY
    run.font.size = Pt(11)
    for item in items:
        p = cell.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.color.rgb = DARK_GREY
        r.font.size = Pt(10.5)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.54)
        section.right_margin  = Cm(2.54)

    # ── Default Normal style ─────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name      = "Calibri"
    normal.font.size      = Pt(10.5)
    normal.font.color.rgb = DARK_GREY

    # ── Header ───────────────────────────────────────────────────────────────
    header      = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text      = "GPS Tracking Solution  |  Proposal"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run            = header_para.runs[0]
    header_run.font.size      = Pt(9)
    header_run.font.color.rgb = TEAL
    header_run.font.italic    = True

    # ── Footer ───────────────────────────────────────────────────────────────
    footer      = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text      = f"Confidential  ·  Generated {datetime.date.today().strftime('%B %d, %Y')}"
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run            = footer_para.runs[0]
    footer_run.font.size      = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    footer_run.font.italic    = True

    # ════════════════════════════════════════════════════════════════════════
    # COVER / TITLE BLOCK
    # ════════════════════════════════════════════════════════════════════════
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("GPS Tracking Solution Proposal")
    title_run.bold           = True
    title_run.font.size      = Pt(26)
    title_run.font.color.rgb = NAVY
    title_run.font.name      = "Calibri"

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Fleet Visibility • Operational Control • Business Resilience")
    sub_run.font.size      = Pt(14)
    sub_run.font.color.rgb = TEAL
    sub_run.font.name      = "Calibri"
    subtitle_p.paragraph_format.space_after = Pt(18)

    add_horizontal_rule(doc)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"
    meta_table.autofit = True
    meta_data = [
        ("Prepared For", "Client Organization"),
        ("Prepared By", "Fleet Management Solutions Team"),
        ("Document Type", "GPS Tracking Solution Proposal"),
        ("Date", datetime.date.today().strftime("%B %d, %Y")),
    ]

    for i, (label, value) in enumerate(meta_data):
        row = meta_table.rows[i]
        label_cell = row.cells[0]
        set_cell_bg(label_cell, "1A376C")
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp  = label_cell.paragraphs[0]
        lr  = lp.add_run(label)
        lr.bold           = True
        lr.font.color.rgb = WHITE
        lr.font.size      = Pt(10.5)
        lr.font.name      = "Calibri"
        lp.paragraph_format.left_indent = Pt(6)

        value_cell = row.cells[1]
        set_cell_bg(value_cell, "E8EEF7")
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        vp  = value_cell.paragraphs[0]
        vr  = vp.add_run(value)
        vr.font.color.rgb = DARK_GREY
        vr.font.size      = Pt(10.5)
        vr.font.name      = "Calibri"
        vp.paragraph_format.left_indent = Pt(6)

    doc.add_paragraph()

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 1 – Executive Summary
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "1.", "Executive Summary")
    add_body(
        doc,
        "This proposal recommends adopting a mature, enterprise-grade GPS tracking platform "
        "that delivers continuous fleet visibility, tighter operational control, and measurable "
        "cost reduction across vehicles and mobile assets. The solution turns location activity "
        "into actionable oversight for executives and operations leaders, enabling faster "
        "decisions, stronger compliance, and improved service reliability."
    )
    add_body(
        doc,
        "With secure access, role-based oversight, and reliable reporting in one environment, "
        "the organization gains driver accountability, asset protection, and clearer performance "
        "management. The platform scales from pilot to nationwide operations without disruption, "
        "protecting business continuity while delivering a rapid return on investment."
    )

    add_highlight_box(
        doc,
        "Strategic Outcomes",
        [
            "Reduce operating costs through tighter trip control and utilization visibility.",
            "Improve driver accountability with policy-based alerts and audit trails.",
            "Strengthen service quality with accurate ETAs and route verification.",
            "Scale confidently across departments, regions, and asset classes.",
        ],
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 2 – Business Challenges Addressed
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "2.", "Business Challenges Addressed")
    add_body(
        doc,
        "Fleet and field operations often face visibility gaps, rising costs, and inconsistent "
        "control across dispersed assets. This proposal addresses the most common barriers to "
        "operational excellence:"
    )
    for label, body in [
        ("Limited real-time visibility:", "Operations teams lack a trusted, live view of fleet status."),
        ("Escalating operating costs:", "Fuel waste, idle time, and unauthorized use go undetected."),
        ("Inconsistent driver accountability:", "Behavior issues are difficult to measure and correct."),
        ("Asset security exposure:", "Theft, misuse, and unplanned downtime threaten continuity."),
        ("Service quality pressure:", "ETA accuracy and delivery reliability are hard to sustain."),
        ("Fragmented reporting:", "Manual logs and disconnected data reduce decision speed."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 3 – Expected Business Benefits
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "3.", "Expected Business Benefits")
    add_body(
        doc,
        "Adopting this GPS tracking platform delivers measurable outcomes for executives, "
        "operations managers, and frontline teams."
    )
    for label, body in [
        ("Fleet visibility:", "A single source of truth for asset location, status, and exceptions."),
        ("Operational control:", "Faster response to deviations, delays, and risk events."),
        ("Cost reduction:", "Lower fuel consumption, overtime, and unplanned maintenance spend."),
        ("Driver accountability:", "Clear, defensible records that support coaching and policy adherence."),
        ("Asset security:", "Improved deterrence and rapid recovery through controlled monitoring."),
        ("Service quality improvement:", "More accurate ETAs, fewer missed deliveries, and higher trust."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 4 – Solution Overview
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "4.", "Solution Overview")
    add_body(
        doc,
        "The solution provides a unified operational control environment that consolidates "
        "tracking data from dedicated devices and approved smartphones. It supports multi-"
        "department structures, role-based oversight, and policy-driven alerts so managers can "
        "proactively manage exceptions rather than react to incidents."
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 5 – Platform Components
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "5.", "Platform Components")

    add_sub_heading(doc, "5.1 Core Tracking & Data Management")
    add_body(
        doc,
        "A resilient operational backbone that captures location updates, standardizes data, "
        "and maintains a trusted record for compliance and decision-making."
    )
    for label, body in [
        ("Continuous location capture:", "Maintains a live operational picture across the fleet."),
        ("Device flexibility:", "Supports dedicated trackers and approved mobile devices."),
        ("Event intelligence:", "Flags overspeeding, boundary breaches, and tamper events."),
        ("Authorized remote actions:", "Supports actions such as immobilization and configuration updates."),
        ("Historical accountability:", "Preserves records for audits, claims, and performance reviews."),
    ]:
        add_bullet(doc, label, body)

    add_sub_heading(doc, "5.2 Web Operations Portal")
    add_body(
        doc,
        "A secure web console for administrators and dispatch teams to monitor fleets, "
        "enforce policies, and generate executive reports from any location."
    )
    for label, body in [
        ("Live fleet overview:", "See real-time locations, status, and exception alerts."),
        ("User and organizational management:", "Role-based access, grouping, and permissions."),
        ("Geofence management:", "Define zones and receive immediate entry/exit alerts."),
        ("Route history and playback:", "Review trips, stops, and idle time for accountability."),
        ("Executive reporting:", "Export trip, mileage, fuel, and utilization summaries."),
    ]:
        add_bullet(doc, label, body)

    add_sub_heading(doc, "5.3 Manager Mobile Application")
    add_body(
        doc,
        "A dedicated mobile experience for supervisors to monitor operations, respond to alerts, "
        "and stay informed while away from the office."
    )
    for label, body in [
        ("Real-time fleet view:", "Access live locations, status filters, and quick summaries."),
        ("Alert center:", "Receive critical event notifications and acknowledge resolution."),
        ("On-the-go reporting:", "Review trips, stop reports, and driver performance."),
        ("Geofence oversight:", "Review zone activity and manage perimeter updates."),
        ("Secure access:", "Protected logins and session controls for field managers."),
    ]:
        add_bullet(doc, label, body)

    add_sub_heading(doc, "5.4 Smartphone Tracking Application")
    add_body(
        doc,
        "A lightweight tracking application that enables approved smartphones to report location, "
        "ideal for contractors, temporary assets, or vehicles without dedicated hardware."
    )
    for label, body in [
        ("Always-on location capture:", "Runs continuously with intelligent battery optimization."),
        ("Offline continuity:", "Stores location points and synchronizes automatically when online."),
        ("Driver status updates:", "Quick status toggles for available, on-trip, or idle."),
        ("Emergency assistance:", "Immediate SOS alerts sent to the control center."),
        ("Trip context notes:", "Annotate deliveries or assignments for better reporting."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 6 – Core Capabilities
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "6.", "Core Capabilities")
    add_body(
        doc,
        "The platform includes a comprehensive set of capabilities that ensure full visibility "
        "and control over fleet operations."
    )
    for label, body in [
        ("Live fleet visibility:", "Up-to-date locations, status, and exceptions for every asset."),
        ("Policy-based alerts:", "Notifications for boundary breaches, deviations, and idle time."),
        ("Driver behavior insights:", "Overspeeding, harsh events, and idling trends for coaching."),
        ("Maintenance control:", "Service schedules aligned to usage to reduce downtime."),
        ("Multi-asset coverage:", "Vehicles, generators, containers, and movable equipment."),
        ("Enterprise integration options:", "Connects with dispatch, ERP, and analytics systems."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 7 – Reporting & Analytics
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "7.", "Reporting & Analytics")
    add_body(
        doc,
        "A comprehensive reporting suite delivers insight into performance, cost, and "
        "compliance to support data-driven decisions."
    )
    for label, body in [
        ("Trip and stop accountability:", "Detailed routes, stops, idle time, and distance metrics."),
        ("Fuel and utilization trends:", "Track consumption patterns and asset utilization rates."),
        ("Incident and compliance trails:", "Downloadable records for investigations and audits."),
        ("Executive dashboards:", "Configurable views for KPIs and operational alerts."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 8 – Reliability & Availability
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "8.", "Reliability & Availability")
    add_body(
        doc,
        "Reliability, availability, and security are built into every layer to keep operations "
        "running without disruption."
    )
    for label, body in [
        ("Continuous operations:", "Redundant services keep tracking online during disruptions."),
        ("Resilient data protection:", "Automated backups and restore procedures safeguard history."),
        ("Offline continuity:", "Devices and apps store data when offline and sync automatically."),
        ("Proactive monitoring:", "System health monitoring and alerts enable rapid response."),
        ("Security governance:", "Role-based access and audit trails protect sensitive data."),
        ("Business continuity options:", "Deployment models and failover planning align to risk needs."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 9 – Implementation & Support
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "9.", "Implementation & Support")
    add_body(
        doc,
        "We provide a structured deployment plan and ongoing support to ensure rapid "
        "adoption, measurable outcomes, and long-term success."
    )
    for label, body in [
        ("Discovery and alignment:", "Confirm fleet structure, policies, and reporting priorities."),
        ("System configuration:", "Configure access governance, alerts, and integrations."),
        ("Onboarding and training:", "Guided sessions for administrators and managers."),
        ("Go-live assurance:", "Live monitoring during launch with rapid issue resolution."),
        ("Managed services:", "Optional monitoring, backups, and enhancements."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 10 – Brand Alignment & Ownership
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "10.", "Brand Alignment & Ownership")
    add_body(
        doc,
        "The platform can be aligned with your corporate identity to ensure a consistent "
        "experience for internal teams and external stakeholders."
    )
    for label, body in [
        ("Custom brand alignment:", "Your logo, colors, and naming across web and mobile apps."),
        ("Branded communications:", "Consistent login screens, email templates, and reports."),
        ("Ownership-ready deliverables:", "Admin access, deployment documentation, and data control."),
    ]:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 11 – Why Organizations Choose This Solution
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "11.", "Why Organizations Choose This Solution")
    add_body(
        doc,
        "Organizations choose this solution because it delivers proven outcomes with the "
        "governance, reliability, and scalability required for enterprise and public-sector use."
    )
    for label, body in [
        ("Enterprise-ready scalability:", "Supports growth from a pilot to nationwide operations."),
        ("Operational control without disruption:", "Improves oversight while keeping teams productive."),
        ("Security and accountability:", "Clear access governance and audit trails for compliance."),
        ("Business continuity focus:", "Resilient operations protect service delivery and reputation."),
        ("Executive visibility:", "Decision-ready reporting for leadership and stakeholders."),
    ]:
        add_bullet(doc, label, body)
    add_body(
        doc,
        "We recommend proceeding with a short discovery and pilot to validate cost savings, "
        "service improvements, and risk reduction in your operating environment. This is a "
        "strategic investment in visibility, control, and resilience, and we are prepared to "
        "partner with your team to deliver measurable ROI and a successful rollout."
    )

    doc.add_paragraph()
    add_horizontal_rule(doc)
    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = closing.add_run("End of Proposal")
    cr.italic           = True
    cr.font.size        = Pt(9)
    cr.font.color.rgb   = RGBColor(0x88, 0x88, 0x88)

    return doc


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    output_path = "GPS_Tracking_Proposal.docx"
    doc = build_document()
    doc.save(output_path)
    print(f"✅  Document saved: {output_path}")
