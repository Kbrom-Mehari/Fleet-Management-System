"""
Generate a well-designed Pricing & Payment Terms Document (.docx)
for the Fleet Management System project.

Requirements
------------
    pip install python-docx

Usage
-----
    python generate_pricing_docx.py

Output
------
    Pricing_and_Payment_Terms.docx  (created in the current working directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


# ---------------------------------------------------------------------------
# Color palette  (identical to generate_sow_docx.py)
# ---------------------------------------------------------------------------
NAVY       = RGBColor(0x1A, 0x37, 0x6C)   # primary heading colour
TEAL       = RGBColor(0x00, 0x7B, 0x83)   # secondary / accent
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_GREY  = RGBColor(0x40, 0x40, 0x40)


# ---------------------------------------------------------------------------
# Helper utilities  (identical styling to generate_sow_docx.py)
# ---------------------------------------------------------------------------

def set_cell_bg(cell, hex_color: str) -> None:
    """Fill a table cell background with a solid color (hex, e.g. '1A376C')."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc: Document) -> None:
    """Insert a thin teal horizontal rule paragraph."""
    p      = doc.add_paragraph()
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "007B83")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(0)


def add_section_heading(doc: Document, number: str, title: str) -> None:
    """Add a numbered section heading (Heading 1 style, navy background strip)."""
    p   = doc.add_paragraph(style="Heading 1")
    run = p.add_run(f"  {number}  {title}")
    run.font.color.rgb = WHITE
    run.font.bold      = True
    run.font.size      = Pt(13)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "1A376C")
    pPr.append(shd)


def add_sub_heading(doc: Document, title: str) -> None:
    """Add a sub-section heading (Heading 2 style, teal text)."""
    p   = doc.add_paragraph(style="Heading 2")
    run = p.add_run(title)
    run.font.color.rgb = TEAL
    run.font.bold      = True
    run.font.size      = Pt(11)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)


def add_body(doc: Document, text: str) -> None:
    """Add a normal body paragraph."""
    p = doc.add_paragraph(text, style="Normal")
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.color.rgb = DARK_GREY
        run.font.size      = Pt(10.5)


def add_bullet(doc: Document, bold_label: str, body_text: str) -> None:
    """Add a bullet point with an optional bold label followed by body text."""
    p         = doc.add_paragraph(style="List Bullet")
    label_run = p.add_run(bold_label)
    label_run.bold           = True
    label_run.font.color.rgb = NAVY
    label_run.font.size      = Pt(10.5)
    if body_text:
        body_run               = p.add_run(" " + body_text)
        body_run.font.color.rgb = DARK_GREY
        body_run.font.size      = Pt(10.5)
    p.paragraph_format.space_after = Pt(4)


def add_note_heading(doc: Document, title: str) -> None:
    """Add a callout / note heading (bold-italic teal, used for sub-phases)."""
    p   = doc.add_paragraph(style="Normal")
    run = p.add_run(title)
    run.bold           = True
    run.italic         = True
    run.font.color.rgb = TEAL
    run.font.size      = Pt(11)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)


# ---------------------------------------------------------------------------
# Pricing table helper
# ---------------------------------------------------------------------------

def add_pricing_table(doc: Document, headers: list, rows: list) -> None:
    """
    Render a styled pricing table.

    headers – list of column header strings
    rows    – list of row tuples  (each tuple has len == len(headers))
    """
    num_cols = len(headers)
    table    = doc.add_table(rows=1 + len(rows), cols=num_cols)
    table.style = "Table Grid"

    # Header row
    hdr_row = table.rows[0]
    for col_idx, hdr_text in enumerate(headers):
        cell = hdr_row.cells[col_idx]
        set_cell_bg(cell, "1A376C")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        run  = para.add_run(hdr_text)
        run.bold           = True
        run.font.color.rgb = WHITE
        run.font.size      = Pt(10.5)
        run.font.name      = "Calibri"
        para.paragraph_format.left_indent = Pt(4)

    # Data rows — alternate background shading
    for row_idx, row_data in enumerate(rows):
        bg    = "E8EEF7" if row_idx % 2 == 0 else "F2F2F2"
        t_row = table.rows[1 + row_idx]
        for col_idx, cell_text in enumerate(row_data):
            cell = t_row.cells[col_idx]
            set_cell_bg(cell, bg)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            run  = para.add_run(cell_text)
            run.font.color.rgb = DARK_GREY
            run.font.size      = Pt(10.5)
            run.font.name      = "Calibri"
            para.paragraph_format.left_indent = Pt(4)

    doc.add_paragraph()   # spacer after table


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.54)
        section.right_margin  = Cm(2.54)

    # ── Default Normal style ─────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name      = "Calibri"
    normal.font.size      = Pt(10.5)
    normal.font.color.rgb = DARK_GREY

    # ── Header ───────────────────────────────────────────────────────────────
    header      = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text      = "Fleet Management System  |  Pricing & Payment Terms"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run            = header_para.runs[0]
    header_run.font.size      = Pt(9)
    header_run.font.color.rgb = TEAL
    header_run.font.italic    = True

    # ── Footer ───────────────────────────────────────────────────────────────
    footer      = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text      = (
        f"Confidential  ·  Generated {datetime.date.today().strftime('%B %d, %Y')}"
    )
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run            = footer_para.runs[0]
    footer_run.font.size      = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    footer_run.font.italic    = True

    # ════════════════════════════════════════════════════════════════════════
    # COVER / TITLE BLOCK
    # ════════════════════════════════════════════════════════════════════════
    title_p           = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run         = title_p.add_run("Pricing & Payment Terms")
    title_run.bold           = True
    title_run.font.size      = Pt(26)
    title_run.font.color.rgb = NAVY
    title_run.font.name      = "Calibri"

    subtitle_p           = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run              = subtitle_p.add_run("Commercial Terms Document")
    sub_run.font.size      = Pt(16)
    sub_run.font.color.rgb = TEAL
    sub_run.font.name      = "Calibri"
    subtitle_p.paragraph_format.space_after = Pt(18)

    add_horizontal_rule(doc)

    # ── Metadata table ───────────────────────────────────────────────────────
    meta_table         = doc.add_table(rows=3, cols=2)
    meta_table.style   = "Table Grid"
    meta_table.autofit = True

    meta_data = [
        ("Project Name",    "Fleet Management System"),
        ("Document Type",   "Pricing & Payment Terms"),
        ("Date",            datetime.date.today().strftime("%B %d, %Y")),
    ]

    for i, (label, value) in enumerate(meta_data):
        row = meta_table.rows[i]

        label_cell = row.cells[0]
        set_cell_bg(label_cell, "1A376C")
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp = label_cell.paragraphs[0]
        lr = lp.add_run(label)
        lr.bold           = True
        lr.font.color.rgb = WHITE
        lr.font.size      = Pt(10.5)
        lr.font.name      = "Calibri"
        lp.paragraph_format.left_indent = Pt(6)

        value_cell = row.cells[1]
        set_cell_bg(value_cell, "E8EEF7")
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        vp = value_cell.paragraphs[0]
        vr = vp.add_run(value)
        vr.font.color.rgb = DARK_GREY
        vr.font.size      = Pt(10.5)
        vr.font.name      = "Calibri"
        vp.paragraph_format.left_indent = Pt(6)

    doc.add_paragraph()   # spacer

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 1 – Itemized Pricing
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "1.", "Itemized Pricing")

    add_body(doc,
        "The following table outlines the itemized cost breakdown for each component of the "
        "Fleet Management System. All prices are in USD and are subject to the final agreed "
        "scope. Placeholder values are marked with $X,XXX and should be updated to reflect "
        "the negotiated amounts before the document is issued to the client."
    )

    # 1a – Licensing
    add_sub_heading(doc, "1.1  Platform Licensing")
    add_body(doc,
        "Licensing covers the right to deploy and operate the platform, including any "
        "proprietary modules, integrations, and the core tracking engine."
    )
    add_pricing_table(doc,
        headers=["License Component", "Description", "Unit Price", "Qty", "Total"],
        rows=[
            ("Core Platform License",    "Annual license for the tracking engine",   "$X,XXX", "1",  "$X,XXX"),
            ("User Seat License",         "Per-seat license for dashboard access",    "$XXX",   "XX", "$X,XXX"),
            ("Mobile App License",        "iOS & Android companion app access",       "$XXX",   "1",  "$XXX"),
            ("API Access License",        "Third-party API integration rights",       "$XXX",   "1",  "$XXX"),
        ]
    )

    # 1b – Setup
    add_sub_heading(doc, "1.2  Setup & Deployment")
    add_body(doc,
        "One-time setup fees cover server provisioning, initial system configuration, "
        "device onboarding, and go-live validation."
    )
    add_pricing_table(doc,
        headers=["Setup Item", "Description", "Cost"],
        rows=[
            ("Server Provisioning & Configuration", "Cloud/on-premise environment setup",               "$X,XXX"),
            ("Database Setup & Optimization",       "Schema creation, indexing, and tuning",            "$X,XXX"),
            ("Device Onboarding (per device)",      "Registration and connectivity verification",       "$XX"),
            ("SSL Certificate Installation",         "HTTPS setup and certificate management",          "$XXX"),
            ("Go-Live Testing & Validation",         "End-to-end pre-production acceptance testing",    "$X,XXX"),
        ]
    )

    # 1c – Customization
    add_sub_heading(doc, "1.3  Customization")
    add_body(doc,
        "Customization fees apply to bespoke development work such as white-labeling, "
        "custom reports, workflow automation, and any feature development outside the "
        "standard platform scope."
    )
    add_pricing_table(doc,
        headers=["Customization Item", "Description", "Estimated Hours", "Rate/Hour", "Estimated Cost"],
        rows=[
            ("White-Labeling & Branding",    "Custom logo, color scheme, and domain",          "XX hrs", "$XXX", "$X,XXX"),
            ("Custom Report Development",    "Tailored analytics dashboards and exports",       "XX hrs", "$XXX", "$X,XXX"),
            ("Workflow Automation",          "Rule-based alerts and automated actions",         "XX hrs", "$XXX", "$X,XXX"),
            ("Third-Party Integration",      "ERP, CRM, or external API integrations",         "XX hrs", "$XXX", "$X,XXX"),
            ("Additional Feature Development","Any scope additions agreed post-contract",       "TBD",    "$XXX", "TBD"),
        ]
    )

    # 1d – Support & Maintenance
    add_sub_heading(doc, "1.4  Support & Maintenance")
    add_body(doc,
        "Ongoing support and maintenance packages ensure the platform remains secure, "
        "up-to-date, and performant throughout the contract period."
    )
    add_pricing_table(doc,
        headers=["Support Tier", "Included Services", "Monthly Cost", "Annual Cost"],
        rows=[
            ("Basic",    "Email support (48-hr response), monthly updates",                 "$XXX",   "$X,XXX"),
            ("Standard", "Priority email & chat (24-hr response), bi-weekly updates",      "$X,XXX", "$XX,XXX"),
            ("Premium",  "24/7 phone/chat, dedicated account manager, weekly updates",     "$X,XXX", "$XX,XXX"),
        ]
    )
    add_body(doc,
        "Note: The hypercare period (first 14–30 days post-launch) is included at no "
        "additional charge under all support tiers."
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 2 – Payment Schedule
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "2.", "Payment Schedule")

    add_body(doc,
        "Payments are structured across project milestones to align financial commitments "
        "with deliverable acceptance. The client may choose one of the following payment "
        "structures at the time of contract signing:"
    )

    # Option A – Milestone-based
    add_note_heading(doc, "Option A: Milestone-Based Payments")
    add_body(doc,
        "Invoices are issued upon the successful completion and client acceptance of each "
        "project milestone:"
    )
    milestone_rows = [
        ("1", "Contract Signing",                    "Project commencement",             "30%",   "$X,XXX"),
        ("2", "Environment Setup & Configuration",   "Server & database live",           "20%",   "$X,XXX"),
        ("3", "System Configuration & UAT",          "User acceptance testing complete",  "25%",   "$X,XXX"),
        ("4", "Go-Live & Handover",                  "Full system handover",             "25%",   "$X,XXX"),
    ]
    add_pricing_table(doc,
        headers=["#", "Milestone", "Trigger", "% of Total", "Amount Due"],
        rows=milestone_rows
    )

    # Option B – Installments
    add_note_heading(doc, "Option B: Monthly Installments")
    add_body(doc,
        "The total project cost is divided into equal monthly installments over an agreed "
        "term. A deposit is required upfront:"
    )
    add_pricing_table(doc,
        headers=["Payment", "Due Date", "Amount"],
        rows=[
            ("Initial Deposit (30%)", "Upon contract signing",   "$X,XXX"),
            ("Installment 1 of X",    "Month 1 after signing",   "$X,XXX"),
            ("Installment 2 of X",    "Month 2 after signing",   "$X,XXX"),
            ("Installment X of X",    "Final month of term",     "$X,XXX"),
        ]
    )

    # Option C – Upfront
    add_note_heading(doc, "Option C: Full Upfront Payment")
    add_body(doc,
        "A 5% discount is applied to the total project cost when full payment is received "
        "within 5 business days of contract signing."
    )
    add_pricing_table(doc,
        headers=["Payment", "Due Date", "Amount"],
        rows=[
            ("Full Project Cost",          "Within 5 business days of signing", "$XX,XXX"),
            ("Upfront Discount (5%)",       "Applied automatically",             "– $X,XXX"),
            ("Net Amount Due",              "",                                  "$XX,XXX"),
        ]
    )

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 3 – Payment Methods Accepted
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "3.", "Payment Methods Accepted")

    add_body(doc,
        "The following payment methods are accepted. All transactions must reference the "
        "invoice number to ensure correct allocation. Currency: USD (other currencies "
        "available upon request)."
    )

    payment_methods = [
        (
            "Bank Wire Transfer (ACH/SWIFT):",
            "Preferred method for large or international transactions. Banking details will "
            "be provided on the invoice. Please allow 3–5 business days for international "
            "transfers to clear.",
        ),
        (
            "Credit / Debit Card:",
            "Visa, Mastercard, and American Express are accepted via our secure payment "
            "portal. A 2.9% processing fee applies to card transactions.",
        ),
        (
            "PayPal / Stripe:",
            "Available for clients preferring online payment platforms. Standard platform "
            "fees may apply.",
        ),
        (
            "Cheque (Check):",
            "Payable to [Company Legal Name]. Please mail to [Company Address]. Cheques "
            "must clear before work commences. Allow 5–7 business days for processing.",
        ),
        (
            "Purchase Order (PO):",
            "Accepted from pre-approved corporate and government clients. A signed PO must "
            "be provided before work begins. Net-30 payment terms apply unless otherwise "
            "agreed in writing.",
        ),
    ]

    for label, body in payment_methods:
        add_bullet(doc, label, body)

    # ════════════════════════════════════════════════════════════════════════
    # SECTION 4 – Late Payment Policies
    # ════════════════════════════════════════════════════════════════════════
    add_section_heading(doc, "4.", "Late Payment Policies")

    add_body(doc,
        "To maintain a healthy and transparent commercial relationship, the following late "
        "payment policies are in effect from the invoice due date. All policies apply unless "
        "an alternative arrangement has been agreed to in writing prior to the due date."
    )

    late_policies = [
        (
            "Grace Period:",
            "A grace period of 5 business days is granted after the invoice due date before "
            "any late fees are applied.",
        ),
        (
            "Late Payment Fee:",
            "Invoices that remain unpaid after the grace period will incur a late fee of "
            "1.5% per month (18% per annum) on the outstanding balance, calculated from the "
            "original due date.",
        ),
        (
            "Service Suspension:",
            "If an invoice remains unpaid for more than 30 calendar days after the due date, "
            "access to the platform and all associated services may be suspended without "
            "further notice until the outstanding balance is settled in full.",
        ),
        (
            "Reinstatement Fee:",
            "A reinstatement fee of $XXX will be applied to restore service access after a "
            "suspension caused by non-payment.",
        ),
        (
            "Legal Action & Collection Costs:",
            "Accounts overdue by more than 60 days may be referred to a collections agency "
            "or pursued through legal proceedings. The client shall be liable for all "
            "reasonable costs of collection, including attorney fees.",
        ),
        (
            "Disputed Invoices:",
            "If the client believes an invoice is incorrect, a written dispute must be "
            "submitted within 10 business days of the invoice date. Undisputed portions of "
            "the invoice remain due on the original payment date.",
        ),
    ]

    for label, body in late_policies:
        add_bullet(doc, label, body)

    # ── Final rule ────────────────────────────────────────────────────────────
    doc.add_paragraph()
    add_horizontal_rule(doc)
    closing           = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr                = closing.add_run("End of Pricing & Payment Terms Document")
    cr.italic           = True
    cr.font.size        = Pt(9)
    cr.font.color.rgb   = RGBColor(0x88, 0x88, 0x88)

    return doc


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    output_path = "Pricing_and_Payment_Terms.docx"
    doc = build_document()
    doc.save(output_path)
    print(f"✅  Document saved: {output_path}")
